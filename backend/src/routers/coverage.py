"""HMO / insurance coverage router for RoyanHealth."""
from __future__ import annotations

import datetime as dt
import logging
from typing import Optional

import psycopg2.extras
import io

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field

from src.constants import (
    COVERAGE_TYPES_ALLOWED,
    LIMIT_TYPES_ALLOWED,
    TABLE_COVERAGE_AUTHORIZATIONS,
    TABLE_COVERAGE_RULES,
    TABLE_HMO_PLANS,
    TABLE_HMOS,
    TABLE_INSURANCE_CLAIMS,
    TABLE_PATIENT_INSURANCE,
    TABLE_PATIENTS,
    TABLE_SERVICE_CATALOG,
    TABLE_RADIOLOGY_CATALOG,
    TABLE_LAB_CATALOG,
    TABLE_PAYER_SERVICE_PRICES,
)

CLAIM_STATUS_ALLOWED = {"submitted", "pre_authorised", "approved", "rejected", "paid", "disputed"}
from src.db import audit_event, db
from src.middleware import require_role, verify_jwt
from src.services.coverage import CoverageInput, calculate_coverage, estimate_bill_coverage

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/coverage", tags=["Coverage"])


class HMOCreate(BaseModel):
    name: str = Field(..., min_length=1)
    code: Optional[str] = None
    contact_phone: Optional[str] = None
    contact_email: Optional[str] = None
    notes: Optional[str] = None


class HMOUpdate(BaseModel):
    name: Optional[str] = Field(None, min_length=1)
    code: Optional[str] = None
    contact_phone: Optional[str] = None
    contact_email: Optional[str] = None
    notes: Optional[str] = None
    is_active: Optional[bool] = None


class HMOPlanCreate(BaseModel):
    name: str = Field(..., min_length=1)
    plan_code: Optional[str] = None
    notes: Optional[str] = None


class ServiceCreate(BaseModel):
    service_code: str = Field(..., min_length=1)
    name: str = Field(..., min_length=1)
    category: str = "misc"
    base_price: float = Field(0.0, ge=0)


class ServiceUpdate(BaseModel):
    name: Optional[str] = Field(None, min_length=1)
    category: Optional[str] = None
    base_price: Optional[float] = Field(None, ge=0)
    is_active: Optional[bool] = None


class PatientInsuranceCreate(BaseModel):
    patient_id: int
    hmo_id: int
    plan_id: Optional[int] = None
    membership_id: str = Field(..., min_length=1)
    status: str = "active"
    valid_from: Optional[str] = None
    valid_to: Optional[str] = None
    is_primary: bool = True
    source: str = "synbot"
    notes: Optional[str] = None
    # Flexible bucket for per-HMO identity fields (employer, dependent
    # relationship, etc.) — exact fields vary by HMO and are still pending
    # from the client, so this ships now without forcing a schema change
    # once they arrive.
    identity_data: Optional[dict] = None


class CoverageRuleCreate(BaseModel):
    hmo_id: int
    service_id: int
    plan_id: Optional[int] = None
    coverage_type: str = "none"
    coverage_value: float = Field(0.0, ge=0)
    requires_approval: bool = False
    limit_type: str = "none"
    limit_value: Optional[float] = Field(None, ge=0)
    notes: Optional[str] = None


class CoverageRuleUpdate(BaseModel):
    coverage_type: Optional[str] = None
    coverage_value: Optional[float] = Field(None, ge=0)
    requires_approval: Optional[bool] = None
    limit_type: Optional[str] = None
    limit_value: Optional[float] = Field(None, ge=0)
    is_active: Optional[bool] = None
    notes: Optional[str] = None


class PatientInsuranceUpdate(BaseModel):
    hmo_id: Optional[int] = None
    plan_id: Optional[int] = None
    membership_id: Optional[str] = None
    status: Optional[str] = None
    valid_from: Optional[str] = None
    valid_to: Optional[str] = None
    is_primary: Optional[bool] = None
    notes: Optional[str] = None
    identity_data: Optional[dict] = None


class ClaimUpdate(BaseModel):
    status: Optional[str] = None
    approved_amount: Optional[float] = Field(None, ge=0)
    authorisation_code: Optional[str] = None
    notes: Optional[str] = None


class BillCoverageEstimateRequest(BaseModel):
    patient_id: int
    consultation_fee: float = Field(0.0, ge=0)
    lab_charges: float = Field(0.0, ge=0)
    pharmacy_charges: float = Field(0.0, ge=0)
    ward_charges: float = Field(0.0, ge=0)
    procedure_charges: float = Field(0.0, ge=0)
    miscellaneous: float = Field(0.0, ge=0)


class CoverageQuoteRequest(BaseModel):
    patient_id: int
    service_id: Optional[int] = None
    service_code: Optional[str] = None
    service_name: Optional[str] = None
    category: str = "misc"
    quantity: float = Field(1.0, gt=0)
    unit_price: Optional[float] = Field(None, ge=0)
    amount: Optional[float] = Field(None, ge=0)
    encounter_id: Optional[int] = None


@router.get("/hmos")
def list_hmos(active_only: bool = True, user=Depends(verify_jwt)):
    query = db.table(TABLE_HMOS).select("*").order("name")
    if active_only:
        query = query.eq("is_active", True)
    result = query.execute()
    return {"hmos": result.data or [], "count": len(result.data or [])}


@router.post("/hmos", status_code=201)
def create_hmo(payload: HMOCreate, user=Depends(require_role("admin"))):
    data = payload.model_dump(exclude_none=True)
    try:
        result = db.table(TABLE_HMOS).insert(data).execute()
    except Exception as exc:
        logger.error("HMO create failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to create HMO")
    hmo = (result.data or [data])[0]
    audit_event("hmo_created", {"hmo_id": hmo.get("id"), "name": payload.name}, actor_id=user.get("sub"), subject_type="hmo", subject_id=str(hmo.get("id") or ""))
    return {"ok": True, "hmo": hmo}


@router.patch("/hmos/{hmo_id}")
def update_hmo(hmo_id: int, payload: HMOUpdate, user=Depends(require_role("admin"))):
    existing = db.table(TABLE_HMOS).select("id").eq("id", hmo_id).limit(1).execute()
    if not existing.data:
        raise HTTPException(status_code=404, detail=f"HMO {hmo_id} not found")
    updates = payload.model_dump(exclude_none=True)
    if not updates:
        raise HTTPException(status_code=422, detail="No fields to update")
    try:
        result = db.table(TABLE_HMOS).update(updates).eq("id", hmo_id).execute()
    except Exception as exc:
        logger.error("HMO update failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to update HMO")
    hmo = (result.data or [{"id": hmo_id, **updates}])[0]
    audit_event("hmo_updated", {"hmo_id": hmo_id, "fields": list(updates.keys())}, actor_id=user.get("sub"), subject_type="hmo", subject_id=str(hmo_id))
    return {"ok": True, "hmo": hmo}


@router.get("/hmos/{hmo_id}/plans")
def list_hmo_plans(hmo_id: int, active_only: bool = True, user=Depends(verify_jwt)):
    query = db.table(TABLE_HMO_PLANS).select("*").eq("hmo_id", hmo_id).order("name")
    if active_only:
        query = query.eq("is_active", True)
    result = query.execute()
    return {"plans": result.data or [], "count": len(result.data or [])}


@router.post("/hmos/{hmo_id}/plans", status_code=201)
def create_hmo_plan(hmo_id: int, payload: HMOPlanCreate, user=Depends(require_role("admin"))):
    hmo = db.table(TABLE_HMOS).select("id").eq("id", hmo_id).limit(1).execute()
    if not hmo.data:
        raise HTTPException(status_code=404, detail=f"HMO {hmo_id} not found")
    data = {**payload.model_dump(exclude_none=True), "hmo_id": hmo_id}
    try:
        result = db.table(TABLE_HMO_PLANS).insert(data).execute()
    except Exception as exc:
        logger.error("HMO plan create failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to create HMO plan")
    plan = (result.data or [data])[0]
    audit_event("hmo_plan_created", {"hmo_id": hmo_id, "plan_id": plan.get("id")}, actor_id=user.get("sub"), subject_type="hmo_plan", subject_id=str(plan.get("id") or ""))
    return {"ok": True, "plan": plan}


@router.get("/services")
def list_services(
    category: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    limit: int = Query(500, ge=1, le=2000),
    user=Depends(verify_jwt),
):
    """The service catalogue, with a count of negotiated payer rates per row.

    The count matters because a service can look priced while still quoting
    differently for every HMO — and because a service with no negotiated rates
    quotes its list price to everyone, which is worth being able to see.
    """
    query = db.table(TABLE_SERVICE_CATALOG).select("*").eq("is_active", True).order("name")
    if category:
        query = query.eq("category", category)
    if q:
        query = query.ilike("name", f"%{q.strip()}%")
    result = query.limit(limit).execute()
    services = result.data or []

    # One pass over the override table rather than a query per row.
    try:
        ids = [s["id"] for s in services if s.get("id")]
        if ids:
            rows = (
                db.table(TABLE_PAYER_SERVICE_PRICES)
                .select("service_id,price")
                .in_("service_id", ids)
                .execute()
                .data
                or []
            )
            agg: dict = {}
            for r in rows:
                a = agg.setdefault(r["service_id"], {"n": 0, "min": None, "max": None})
                p = float(r.get("price") or 0)
                a["n"] += 1
                a["min"] = p if a["min"] is None else min(a["min"], p)
                a["max"] = p if a["max"] is None else max(a["max"], p)
            for s_ in services:
                a = agg.get(s_.get("id"))
                s_["payer_rate_count"] = (a or {}).get("n", 0)
                s_["payer_rate_min"] = (a or {}).get("min")
                s_["payer_rate_max"] = (a or {}).get("max")
    except Exception as exc:
        logger.warning("payer rate aggregation failed: %s", exc)
        for s_ in services:
            s_.setdefault("payer_rate_count", 0)

    return {"services": services, "count": len(services)}


@router.get("/services/{service_id}/payer-prices")
def list_service_payer_prices(service_id: int, user=Depends(verify_jwt)):
    """Every HMO's negotiated price for one service, highest first.

    This is the view that was impossible before migration 198: Hope priced per
    payer, migration 188 flattened it to a single charge, and the catalogue
    could only ever show one number.
    """
    try:
        rows = (
            db.table(TABLE_PAYER_SERVICE_PRICES)
            .select("id,hmo_id,price,source,updated_at")
            .eq("service_id", service_id)
            .execute()
            .data
            or []
        )
        if rows:
            hmos = (
                db.table(TABLE_HMOS)
                .select("id,name,code")
                .in_("id", [r["hmo_id"] for r in rows])
                .execute()
                .data
                or []
            )
            by_id = {h["id"]: h for h in hmos}
            for r in rows:
                h = by_id.get(r["hmo_id"]) or {}
                r["hmo_name"] = h.get("name")
                r["hmo_code"] = h.get("code")
        rows.sort(key=lambda r: float(r.get("price") or 0), reverse=True)
        return {"prices": rows, "count": len(rows)}
    except Exception as exc:
        logger.error("payer price listing failed for service %s: %s", service_id, exc)
        return {"prices": [], "count": 0}


@router.post("/services", status_code=201)
def create_service(payload: ServiceCreate, user=Depends(require_role("admin"))):
    data = payload.model_dump()
    try:
        result = db.table(TABLE_SERVICE_CATALOG).insert(data).execute()
    except Exception as exc:
        logger.error("Service create failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to create service")
    service = (result.data or [data])[0]
    audit_event("service_catalog_created", {"service_id": service.get("id"), "code": payload.service_code}, actor_id=user.get("sub"), subject_type="service", subject_id=str(service.get("id") or ""))
    return {"ok": True, "service": service}


@router.patch("/services/{service_id}")
def update_service(service_id: int, payload: ServiceUpdate, user=Depends(require_role("admin"))):
    """Price/edit an existing catalog row — the seeded rows (CONSULT-GEN etc.)
    ship with base_price=0, and there was previously no way to fix that
    short of deleting and re-adding under a now-conflicting service_code."""
    existing = db.table(TABLE_SERVICE_CATALOG).select("id").eq("id", service_id).limit(1).execute()
    if not existing.data:
        raise HTTPException(status_code=404, detail=f"Service {service_id} not found")

    updates = payload.model_dump(exclude_none=True)
    if not updates:
        raise HTTPException(status_code=422, detail="No fields to update")

    try:
        result = db.table(TABLE_SERVICE_CATALOG).update(updates).eq("id", service_id).execute()
    except Exception as exc:
        logger.error("Service update failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to update service")
    service = (result.data or [{"id": service_id, **updates}])[0]
    audit_event("service_catalog_updated", {"service_id": service_id, "fields": list(updates.keys())}, actor_id=user.get("sub"), subject_type="service", subject_id=str(service_id))
    return {"ok": True, "service": service}


@router.post("/patient-insurance", status_code=201)
def create_patient_insurance(payload: PatientInsuranceCreate, user=Depends(verify_jwt)):
    data = payload.model_dump(exclude_none=True)
    try:
        result = db.table(TABLE_PATIENT_INSURANCE).insert(data).execute()
    except Exception as exc:
        logger.error("Patient insurance create failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to create patient insurance")
    policy = (result.data or [data])[0]
    audit_event("patient_insurance_created", {"patient_id": payload.patient_id, "hmo_id": payload.hmo_id}, actor_id=user.get("sub"), subject_type="patient", subject_id=str(payload.patient_id))
    return {"ok": True, "insurance": policy}


@router.get("/patient-insurance")
def list_patient_insurance(patient_id: int = Query(...), user=Depends(verify_jwt)):
    """Full policy history for one patient (not just the active/primary one
    get_active_patient_insurance() uses internally) -- for the patient HMO
    workpanel, which needs to show and let staff change enrollment."""
    result = db.table(TABLE_PATIENT_INSURANCE).select("*").eq("patient_id", patient_id).order("created_at", desc=True).execute()
    rows = result.data or []
    hmo_ids = list({r["hmo_id"] for r in rows if r.get("hmo_id")})
    if hmo_ids:
        hmos = db.table(TABLE_HMOS).select("id,name").in_("id", hmo_ids).execute()
        hmos_by_id = {h["id"]: h["name"] for h in (hmos.data or [])}
        for r in rows:
            r["hmo_name"] = hmos_by_id.get(r.get("hmo_id"))
    return {"policies": rows, "count": len(rows)}


@router.patch("/patient-insurance/{policy_id}")
def update_patient_insurance(policy_id: int, payload: PatientInsuranceUpdate, user=Depends(verify_jwt)):
    existing = db.table(TABLE_PATIENT_INSURANCE).select("id,patient_id").eq("id", policy_id).limit(1).execute()
    if not existing.data:
        raise HTTPException(status_code=404, detail=f"Patient insurance record {policy_id} not found")
    updates = payload.model_dump(exclude_none=True)
    if not updates:
        raise HTTPException(status_code=422, detail="No fields to update")
    try:
        result = db.table(TABLE_PATIENT_INSURANCE).update(updates).eq("id", policy_id).execute()
    except Exception as exc:
        logger.error("Patient insurance update failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to update patient insurance (e.g. another active primary policy already exists for this patient)")
    policy = (result.data or [{"id": policy_id, **updates}])[0]
    audit_event("patient_insurance_updated", {"policy_id": policy_id, "fields": list(updates.keys())}, actor_id=user.get("sub"), subject_type="patient", subject_id=str(existing.data[0].get("patient_id")))
    return {"ok": True, "insurance": policy}


@router.post("/rules", status_code=201)
def create_coverage_rule(payload: CoverageRuleCreate, user=Depends(require_role("admin"))):
    if payload.coverage_type not in COVERAGE_TYPES_ALLOWED:
        raise HTTPException(status_code=422, detail=f"coverage_type must be one of {sorted(COVERAGE_TYPES_ALLOWED)}")
    if payload.limit_type not in LIMIT_TYPES_ALLOWED:
        raise HTTPException(status_code=422, detail=f"limit_type must be one of {sorted(LIMIT_TYPES_ALLOWED)}")
    data = payload.model_dump(exclude_none=True)
    try:
        result = db.table(TABLE_COVERAGE_RULES).insert(data).execute()
    except Exception as exc:
        logger.error("Coverage rule create failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to create coverage rule")
    rule = (result.data or [data])[0]
    audit_event("coverage_rule_created", {"rule_id": rule.get("id"), "hmo_id": payload.hmo_id, "service_id": payload.service_id}, actor_id=user.get("sub"), subject_type="coverage_rule", subject_id=str(rule.get("id") or ""))
    return {"ok": True, "rule": rule}


@router.get("/rules")
def list_coverage_rules(hmo_id: Optional[int] = None, service_id: Optional[int] = None, user=Depends(verify_jwt)):
    query = db.table(TABLE_COVERAGE_RULES).select("*").order("created_at", desc=True).limit(500)
    if hmo_id:
        query = query.eq("hmo_id", hmo_id)
    if service_id:
        query = query.eq("service_id", service_id)
    result = query.execute()
    return {"rules": result.data or [], "count": len(result.data or [])}


@router.patch("/rules/{rule_id}")
def update_coverage_rule(rule_id: int, payload: CoverageRuleUpdate, user=Depends(require_role("admin"))):
    if payload.coverage_type and payload.coverage_type not in COVERAGE_TYPES_ALLOWED:
        raise HTTPException(status_code=422, detail=f"coverage_type must be one of {sorted(COVERAGE_TYPES_ALLOWED)}")
    if payload.limit_type and payload.limit_type not in LIMIT_TYPES_ALLOWED:
        raise HTTPException(status_code=422, detail=f"limit_type must be one of {sorted(LIMIT_TYPES_ALLOWED)}")
    existing = db.table(TABLE_COVERAGE_RULES).select("id").eq("id", rule_id).limit(1).execute()
    if not existing.data:
        raise HTTPException(status_code=404, detail=f"Coverage rule {rule_id} not found")
    updates = payload.model_dump(exclude_none=True)
    if not updates:
        raise HTTPException(status_code=422, detail="No fields to update")
    try:
        result = db.table(TABLE_COVERAGE_RULES).update(updates).eq("id", rule_id).execute()
    except Exception as exc:
        logger.error("Coverage rule update failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to update coverage rule")
    rule = (result.data or [{"id": rule_id, **updates}])[0]
    audit_event("coverage_rule_updated", {"rule_id": rule_id, "fields": list(updates.keys())}, actor_id=user.get("sub"), subject_type="coverage_rule", subject_id=str(rule_id))
    return {"ok": True, "rule": rule}


@router.get("/patients")
def list_covered_patients(
    hmo_id: Optional[int] = Query(None),
    search: Optional[str] = Query(None),
    limit: int = Query(100, ge=1, le=500),
    user=Depends(verify_jwt),
):
    """Patient roster under HMO cover — matches Hope's HMO Dashboard patient
    list (registration date, name, insurance no., age/gender, HMO, status).
    Joins the modern royan_patient_insurance table, falling back to the
    legacy royan_patients.hmo_provider field for patients registered before
    that table existed (same fallback used elsewhere in this app)."""
    conn = db.pool.getconn()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            sql = """
                SELECT p.id AS patient_id, p.mrn, p.first_name, p.last_name,
                       p.date_of_birth, p.gender, p.registered_at, p.patient_type,
                       COALESCE(h.name, p.hmo_provider) AS hmo_name,
                       pi.hmo_id AS hmo_id,
                       COALESCE(pi.membership_id, p.hmo_id::text) AS membership_id,
                       COALESCE(pi.status, 'active') AS status
                FROM public.royan_patients p
                LEFT JOIN LATERAL (
                    SELECT * FROM public.royan_patient_insurance ins
                    WHERE ins.patient_id = p.id AND ins.is_primary = true AND ins.status = 'active'
                    ORDER BY ins.created_at DESC LIMIT 1
                ) pi ON true
                LEFT JOIN public.royan_hmos h ON h.id = pi.hmo_id
                WHERE (pi.id IS NOT NULL OR NULLIF(trim(p.hmo_provider), '') IS NOT NULL)
            """
            params: list = []
            if hmo_id:
                sql += " AND pi.hmo_id = %s"
                params.append(hmo_id)
            if search:
                sql += " AND (p.first_name ILIKE %s OR p.last_name ILIKE %s OR p.mrn ILIKE %s)"
                like = f"%{search.strip()}%"
                params.extend([like, like, like])
            sql += " ORDER BY upper(p.first_name), upper(p.last_name), p.id LIMIT %s"   # A-Z, as the names show
            params.append(limit)
            cur.execute(sql, tuple(params))
            rows = [dict(r) for r in cur.fetchall()]
        conn.commit()
    except Exception as exc:
        try:
            conn.rollback()
        except Exception:
            pass
        logger.error("Covered-patients roster query failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to load HMO patient roster")
    finally:
        db.pool.putconn(conn)

    for r in rows:
        r["patient_name"] = f"{r.pop('first_name', '') or ''} {r.pop('last_name', '') or ''}".strip() or None

    return {"patients": rows, "count": len(rows)}


@router.post("/quote")
def quote_coverage(payload: CoverageQuoteRequest, user=Depends(verify_jwt)):
    quote = calculate_coverage(CoverageInput(**payload.model_dump()))
    return {"ok": True, "quote": quote}


@router.post("/estimate-bill")
def estimate_bill(payload: BillCoverageEstimateRequest, user=Depends(verify_jwt)):
    """Auto-suggest the HMO-covered portion of a flat-category bill (the
    GenerateBillPanel form) so staff see the discount already applied
    instead of typing it in by hand — matches Hope's behavior where an
    HMO-attached patient's bill template already reflects their coverage."""
    result = estimate_bill_coverage(payload.patient_id, {
        "consultation_fee": payload.consultation_fee,
        "lab_charges": payload.lab_charges,
        "pharmacy_charges": payload.pharmacy_charges,
        "ward_charges": payload.ward_charges,
        "procedure_charges": payload.procedure_charges,
        "miscellaneous": payload.miscellaneous,
    })
    return result


# ─── Claims ───────────────────────────────────────────────────────────────────

@router.get("/claims")
def list_claims(
    patient_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    hmo_provider_id: Optional[int] = Query(None),
    limit: int = Query(100, ge=1, le=500),
    user=Depends(verify_jwt),
):
    query = db.table(TABLE_INSURANCE_CLAIMS).select("*").order("submitted_at", desc=True).limit(limit)
    if patient_id:
        query = query.eq("patient_id", patient_id)
    if status:
        query = query.eq("status", status)
    if hmo_provider_id:
        query = query.eq("hmo_provider_id", hmo_provider_id)
    result = query.execute()
    claims = result.data or []

    patient_ids = list({c["patient_id"] for c in claims if c.get("patient_id")})
    patients_by_id: dict[int, dict] = {}
    if patient_ids:
        patients = db.table(TABLE_PATIENTS).select("id,mrn,first_name,last_name").in_("id", patient_ids).execute()
        patients_by_id = {p["id"]: p for p in (patients.data or [])}
    for c in claims:
        p = patients_by_id.get(c.get("patient_id"))
        c["patient_name"] = f"{p['first_name']} {p['last_name']}" if p else None
        c["mrn"] = p.get("mrn") if p else None

    return {"claims": claims, "count": len(claims)}


@router.patch("/claims/{claim_id}")
def update_claim(claim_id: int, payload: ClaimUpdate, user=Depends(require_role("admin", "finance", "hmo"))):
    """Confirm/progress a claim (submitted -> pre_authorised/approved/rejected
    -> paid/disputed). Moving to 'paid' auto-stamps settled_at unless the
    caller already knows a specific settlement date -- that's not exposed
    here since staff confirming a claim are marking "this happened now",
    not backdating one."""
    if payload.status and payload.status not in CLAIM_STATUS_ALLOWED:
        raise HTTPException(status_code=422, detail=f"status must be one of {sorted(CLAIM_STATUS_ALLOWED)}")
    existing = db.table(TABLE_INSURANCE_CLAIMS).select("id").eq("id", claim_id).limit(1).execute()
    if not existing.data:
        raise HTTPException(status_code=404, detail=f"Claim {claim_id} not found")
    updates = payload.model_dump(exclude_none=True)
    if not updates:
        raise HTTPException(status_code=422, detail="No fields to update")
    if updates.get("status") == "paid":
        updates["settled_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    try:
        result = db.table(TABLE_INSURANCE_CLAIMS).update(updates).eq("id", claim_id).execute()
    except Exception as exc:
        logger.error("Claim update failed: %s", exc)
        raise HTTPException(status_code=500, detail="Failed to update claim")
    claim = (result.data or [{"id": claim_id, **updates}])[0]
    audit_event(f"claim_{payload.status or 'updated'}", {"claim_id": claim_id, "fields": list(updates.keys())}, actor_id=user.get("sub"), subject_type="claim", subject_id=str(claim_id))
    return {"ok": True, "claim": claim}


# ─── HMO Authorization Workflow ──────────────────────────────────────────────

class AuthorizationCreate(BaseModel):
    patient_id: int
    service_id: Optional[int] = None
    service_name: Optional[str] = None
    hmo_id: Optional[int] = None
    hmo_name: Optional[str] = None
    requested_by: Optional[str] = None
    notes: Optional[str] = None


class AuthorizationResolve(BaseModel):
    status: str  # "approved" | "rejected"
    approval_code: Optional[str] = None
    notes: Optional[str] = None


@router.post("/authorizations", status_code=201)
def request_authorization(payload: AuthorizationCreate, user=Depends(verify_jwt)):
    """Create a pre-authorization request for a service that requires HMO approval."""
    data = {
        **payload.model_dump(exclude_none=True),
        "status": "requested",
        "requested_by": payload.requested_by or user.get("email") or user.get("sub"),
    }
    result = db.table(TABLE_COVERAGE_AUTHORIZATIONS).insert(data).execute()
    auth = (result.data or [data])[0]
    audit_event("authorization_requested", {"patient_id": payload.patient_id, "service_name": payload.service_name}, actor_id=user.get("sub"), subject_type="authorization", subject_id=str(auth.get("id") or ""))
    return {"ok": True, "authorization": auth}


@router.get("/authorizations")
def list_authorizations(
    status: Optional[str] = None,
    hmo_id: Optional[int] = None,
    limit: int = 100,
    user=Depends(verify_jwt),
):
    query = db.table(TABLE_COVERAGE_AUTHORIZATIONS).select("*").order("requested_at", desc=True).limit(limit)
    if status:
        query = query.eq("status", status)
    if hmo_id:
        query = query.eq("hmo_id", hmo_id)
    result = query.execute()
    authorizations = result.data or []

    patient_ids = list({a["patient_id"] for a in authorizations if a.get("patient_id")})
    patients_by_id: dict[int, dict] = {}
    if patient_ids:
        patients = db.table(TABLE_PATIENTS).select("id,mrn,first_name,last_name").in_("id", patient_ids).execute()
        patients_by_id = {p["id"]: p for p in (patients.data or [])}

    for a in authorizations:
        p = patients_by_id.get(a.get("patient_id"))
        a["patient_name"] = f"{p['first_name']} {p['last_name']}" if p else None
        a["mrn"] = p.get("mrn") if p else None

    return {"authorizations": authorizations, "count": len(authorizations)}


@router.patch("/authorizations/{auth_id}")
def resolve_authorization(auth_id: int, payload: AuthorizationResolve, user=Depends(require_role("admin", "finance", "hmo"))):
    if payload.status not in ("approved", "rejected"):
        raise HTTPException(status_code=422, detail="status must be 'approved' or 'rejected'")
    update = {
        "status": payload.status,
        "decided_by": user.get("email") or user.get("sub"),
    }
    if payload.approval_code:
        update["authorization_code"] = payload.approval_code
    if payload.notes:
        update["notes"] = payload.notes
    result = db.table(TABLE_COVERAGE_AUTHORIZATIONS).update(update).eq("id", auth_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail=f"Authorization {auth_id} not found")
    audit_event(f"authorization_{payload.status}", {"approval_code": payload.approval_code}, actor_id=user.get("sub"), subject_type="authorization", subject_id=str(auth_id))
    return {"ok": True, "authorization": result.data[0]}


# -- HMO tariff spreadsheet import -------------------------------------------
#
# "each HMO provider has an Excel file that shows all the tariffs for all the
# services that HMO provides with the pricing... so that if there's a tariff
# change, we can just upload the new Excel file and update the HMO's details."
#
# TWO PHASES ON PURPOSE
# ---------------------
# An import writes negotiated prices for one payer across hundreds of services.
# Get a column wrong and every price for that HMO is wrong, silently, until an
# invoice is disputed. So the upload only ever RETURNS a preview; a separate
# call commits rows the user has seen. Nothing is written by the act of
# uploading.

_TARIFF_NAME_HEADERS = {"service", "service name", "description", "item", "test",
                        "procedure", "name", "tariff", "benefit"}
_TARIFF_PRICE_HEADERS = {"price", "amount", "tariff", "rate", "cost", "fee",
                         "charge", "ngn", "naira", "value"}


def _norm(v) -> str:
    return " ".join(str(v or "").strip().lower().split())


def _parse_tariff_sheet(content: bytes) -> Dict[str, Any]:
    """Pull (service, price) pairs out of a spreadsheet with unknown columns.

    Every HMO formats differently, so the header row is located by looking for
    a row containing both a name-like and a price-like heading rather than
    assuming row 1. Rows above it (titles, logos, blank spacers) are skipped.
    """
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - openpyxl is in requirements
        raise HTTPException(status_code=500, detail="Spreadsheet support unavailable")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(
            status_code=422,
            detail="Could not read that file. Save it as .xlsx and try again.",
        )

    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=422, detail="That spreadsheet is empty.")

    name_col = price_col = header_row = None
    for i, row in enumerate(rows[:25]):
        cells = [_norm(c) for c in row]
        # Substring, not equality. Real sheets head this column "Service
        # Description", "Name of Service", "Test/Procedure" -- an exact match
        # against a fixed list rejects nearly every one of them.
        n_idx = next((j for j, c in enumerate(cells)
                      if c and any(k in c for k in _TARIFF_NAME_HEADERS)), None)
        p_idx = next((j for j, c in enumerate(cells)
                      if c and j != n_idx
                      and any(k in c for k in _TARIFF_PRICE_HEADERS)), None)
        if n_idx is not None and p_idx is not None and n_idx != p_idx:
            name_col, price_col, header_row = n_idx, p_idx, i
            break

    if header_row is None:
        raise HTTPException(
            status_code=422,
            detail=("Could not find the service and price columns. The sheet needs a "
                    "header row with a service/description column and a price/amount "
                    "column."),
        )

    parsed, skipped = [], 0
    for row in rows[header_row + 1:]:
        if name_col >= len(row) or price_col >= len(row):
            continue
        name = str(row[name_col] or "").strip()
        raw = row[price_col]
        if not name:
            continue
        try:
            price = float(str(raw).replace(",", "").replace("N", "").replace("\u20a6", "").strip())
        except (TypeError, ValueError):
            skipped += 1
            continue
        if price <= 0:
            skipped += 1
            continue
        parsed.append({"name": name, "price": price})

    return {"rows": parsed, "unreadable": skipped,
            "header_row": header_row + 1,
            "columns": {"service": name_col + 1, "price": price_col + 1}}


def _match_tariff_rows(rows: List[Dict]) -> List[Dict]:
    """Attach each spreadsheet row to a catalogue entry, or mark it unmatched.

    Matched against all three catalogues because an HMO tariff sheet mixes
    services, lab tests and imaging in one list. Exact name only -- fuzzy
    matching a tariff line is how a cheap test acquires an expensive test's
    price.
    """
    lookup: Dict[str, Dict[str, Any]] = {}
    for table, kind, col in (
        (TABLE_SERVICE_CATALOG, "service", "service_id"),
        (TABLE_LAB_CATALOG, "lab", "lab_catalog_id"),
        (TABLE_RADIOLOGY_CATALOG, "radiology", "radiology_catalog_id"),
    ):
        try:
            for r in db.table(table).select("id,name").execute().data or []:
                key = _norm(r.get("name"))
                # First catalogue wins: a name in two catalogues is ambiguous,
                # and silently preferring the last one read would be arbitrary.
                if key and key not in lookup:
                    lookup[key] = {"id": r["id"], "kind": kind, "column": col,
                                   "matched_name": r.get("name")}
        except Exception as exc:
            logger.warning("Tariff import: could not read %s: %s", table, exc)

    out = []
    for r in rows:
        hit = lookup.get(_norm(r["name"]))
        out.append({**r, "matched": bool(hit), **(hit or {})})
    return out


@router.post("/hmos/{hmo_id}/tariff-import/preview")
async def preview_tariff_import(
    hmo_id: int,
    file: UploadFile = File(...),
    user=Depends(require_role("admin")),
):
    """Read a tariff spreadsheet and show what WOULD change. Writes nothing."""
    hmo = db.table(TABLE_HMOS).select("id,name").eq("id", hmo_id).limit(1).execute()
    if not hmo.data:
        raise HTTPException(status_code=404, detail=f"HMO {hmo_id} not found")

    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="That file is larger than 5 MB.")

    parsed = _parse_tariff_sheet(content)
    matched = _match_tariff_rows(parsed["rows"])

    # What each price would replace, so the reviewer sees the delta rather than
    # just the incoming number.
    existing: Dict[Any, float] = {}
    try:
        cur = db.table(TABLE_PAYER_SERVICE_PRICES).select(
            "service_id,lab_catalog_id,radiology_catalog_id,price"
        ).eq("hmo_id", hmo_id).execute().data or []
        for c in cur:
            for col in ("service_id", "lab_catalog_id", "radiology_catalog_id"):
                if c.get(col):
                    existing[(col, c[col])] = float(c.get("price") or 0)
    except Exception as exc:
        logger.warning("Tariff import: existing price lookup failed: %s", exc)

    for m in matched:
        if m.get("matched"):
            m["current_price"] = existing.get((m["column"], m["id"]))

    ok = [m for m in matched if m.get("matched")]
    return {
        "hmo": hmo.data[0],
        "filename": file.filename,
        "parsed_from_row": parsed["header_row"],
        "columns": parsed["columns"],
        "total_rows": len(matched),
        "matched": len(ok),
        "unmatched": len(matched) - len(ok),
        "unreadable_price": parsed["unreadable"],
        "changed": sum(1 for m in ok
                       if m.get("current_price") is not None
                       and float(m["current_price"]) != float(m["price"])),
        "new": sum(1 for m in ok if m.get("current_price") is None),
        "rows": matched[:500],
    }


class TariffCommitRow(BaseModel):
    id: int
    column: str
    price: float


class TariffCommit(BaseModel):
    rows: List[TariffCommitRow]


@router.post("/hmos/{hmo_id}/tariff-import/commit")
def commit_tariff_import(
    hmo_id: int,
    payload: TariffCommit,
    user=Depends(require_role("admin")),
):
    """Write the rows the user reviewed. Only the three catalogue columns are
    accepted, so a crafted payload cannot write to an arbitrary column."""
    allowed = {"service_id", "lab_catalog_id", "radiology_catalog_id"}
    written = failed = 0
    for r in payload.rows:
        if r.column not in allowed:
            failed += 1
            continue
        try:
            db.table(TABLE_PAYER_SERVICE_PRICES).upsert(
                {"hmo_id": hmo_id, r.column: r.id, "price": r.price,
                 "source": "tariff_upload"},
                on_conflict="hmo_id,%s" % r.column,
            ).execute()
            written += 1
        except Exception as exc:
            logger.warning("Tariff commit failed for %s=%s: %s", r.column, r.id, exc)
            failed += 1

    audit_event("hmo_tariff_imported",
                {"hmo_id": hmo_id, "written": written, "failed": failed},
                actor_id=user.get("sub"), subject_type="hmo", subject_id=str(hmo_id))
    invalidate_cache_tags("coverage", "hospital")
    return {"ok": True, "written": written, "failed": failed}


# ─── HMO desk: the HMO patients on today's queue ──────────────────────────────
# The HMO officer's own queue (finance meeting, Oct 2026): only the HMO patients, where each one is in the visit,
# what has been ordered and dispensed for them (the pharmacy's sign-off shows here as it happens), the bills and the
# HMO's share, the patient's signature on the claims form, and the officer's check before the claim goes out.

_HMO_DESK_ENCOUNTERS = """
    SELECT e.id, e.queue_number, e.status, e.encounter_type, e.created_at, e.discharge_at, e.consulting_doctor,
           e.chief_complaint, e.claim_signature_url, e.claim_signed_at, e.claim_signed_by,
           e.hmo_verified_at, e.hmo_verified_by, e.hmo_verify_note,
           p.id AS patient_id, p.mrn, p.first_name, p.last_name, p.gender, p.date_of_birth,
           COALESCE(h.name, NULLIF(trim(p.hmo_provider), '')) AS hmo_name, pi.hmo_id,
           COALESCE(pi.membership_id, e.membership_id, p.hmo_id::text) AS membership_id
    FROM royan_encounters e
    JOIN royan_patients p ON p.id = e.patient_id
    LEFT JOIN LATERAL (SELECT * FROM royan_patient_insurance i WHERE i.patient_id = p.id AND i.status = 'active'
                       ORDER BY i.is_primary DESC, i.id LIMIT 1) pi ON true
    LEFT JOIN royan_hmos h ON h.id = pi.hmo_id
    WHERE e.external_source IS NULL   -- the live queue only, never Hope's imported visits
      AND (e.payer_type = 'hmo' OR (e.payer_type IS NULL AND (pi.id IS NOT NULL OR NULLIF(trim(p.hmo_provider), '') IS NOT NULL)))
      AND (e.status <> 'discharged' OR e.discharge_at >= (date_trunc('day', now() AT TIME ZONE 'Africa/Lagos') - make_interval(days => %s)) AT TIME ZONE 'Africa/Lagos')
      AND (%s::int IS NULL OR e.id = %s::int)
    ORDER BY (e.status = 'discharged'), e.created_at
    LIMIT 300
"""


def _hmo_desk_rows(days: int, encounter_id: Optional[int] = None) -> list[dict]:
    from src.fin.db import q as _q, tx as _tx
    with _tx() as conn:
        encs = _q(conn, _HMO_DESK_ENCOUNTERS, (days, encounter_id, encounter_id))
        ids = [e["id"] for e in encs]
        if not ids:
            return []
        rx = _q(conn, """SELECT p.encounter_id, p.id, p.drug_name, p.dosage, p.frequency, p.duration, p.quantity, p.status,
                                p.approval_status, d.quantity_dispensed, d.total_cost, d.dispensed_at, d.dispensed_by,
                                d.signature_url
                         FROM royan_prescriptions p
                         LEFT JOIN LATERAL (SELECT * FROM royan_dispensing x WHERE x.prescription_id = p.id
                                            ORDER BY x.dispensed_at DESC LIMIT 1) d ON true
                         WHERE p.encounter_id = ANY(%s) ORDER BY p.prescribed_at""", (ids,))
        labs = _q(conn, """SELECT encounter_id, test_name, status FROM royan_lab_orders
                           WHERE encounter_id = ANY(%s) ORDER BY ordered_at""", (ids,))
        from src.routers.billing_hospital import visit_diagnoses
        dx = [{**d, "encounter_id": i} for i in ids for d in visit_diagnoses(conn, i)]
        bills = _q(conn, """SELECT b.id, b.encounter_id, b.bill_number, b.subtotal, b.discount, b.hmo_covered,
                                   b.patient_balance, b.status, b.stage, b.billed_at,
                                   COALESCE((SELECT SUM(amount) FROM royan_payments x WHERE x.billing_id = b.id
                                             AND COALESCE(x.status, 'posted') <> 'void'), 0) AS paid,
                                   c.id AS claim_id, c.status AS claim_status, c.claim_amount, c.authorisation_code
                            FROM royan_billing b
                            LEFT JOIN LATERAL (SELECT * FROM royan_insurance_claims c WHERE c.billing_id = b.id
                                               ORDER BY c.id DESC LIMIT 1) c ON true
                            WHERE b.encounter_id = ANY(%s) AND b.status <> 'void' ORDER BY b.id""", (ids,))
        auths = _q(conn, """SELECT encounter_id, status, COUNT(*) AS n FROM royan_coverage_authorizations
                            WHERE encounter_id = ANY(%s) GROUP BY 1, 2""", (ids,))

    def by(rows):
        out: dict = {i: [] for i in ids}
        for r in rows:
            out.setdefault(r["encounter_id"], []).append(r)
        return out

    rx_by, lab_by, dx_by, bill_by = by(rx), by(labs), by(dx), by(bills)
    money = ("subtotal", "discount", "hmo_covered", "patient_balance", "paid", "claim_amount")
    out = []
    for e in encs:
        drugs = [{"id": r["id"], "drug": r["drug_name"],
                  "dosage": " ".join(x for x in (r.get("dosage"), r.get("frequency")) if x),
                  "duration": r.get("duration"), "quantity": r.get("quantity_dispensed") or r.get("quantity"),
                  "status": r.get("status"), "approval_status": r.get("approval_status"),
                  "dispensed": bool(r.get("dispensed_at")), "dispensed_at": r.get("dispensed_at"),
                  "dispensed_by": r.get("dispensed_by"), "cost": float(r.get("total_cost") or 0),
                  "signed": bool(r.get("signature_url"))} for r in rx_by[e["id"]]]
        bl = [{**b, **{k: float(b[k] or 0) for k in money}} for b in bill_by[e["id"]]]
        a = {r["status"]: int(r["n"]) for r in auths if r["encounter_id"] == e["id"]}
        out.append({
            "encounter_id": e["id"], "queue_number": e["queue_number"], "status": e["status"],
            "encounter_type": e["encounter_type"], "arrived_at": e["created_at"], "discharged_at": e["discharge_at"],
            "doctor": e["consulting_doctor"], "chief_complaint": e["chief_complaint"],
            "patient_id": e["patient_id"], "mrn": e["mrn"],
            "patient_name": " ".join(x for x in (e.get("first_name"), e.get("last_name")) if x) or e["mrn"],
            "gender": e["gender"], "date_of_birth": e["date_of_birth"],
            "hmo_name": e["hmo_name"], "hmo_id": e["hmo_id"], "membership_id": e["membership_id"],
            "diagnoses": [{"code": d["icd_code"], "description": d["description"], "type": d["diagnosis_type"]}
                          for d in dx_by[e["id"]]],
            "labs": [{"test": l["test_name"], "status": l["status"]} for l in lab_by[e["id"]]],
            "drugs": drugs, "drugs_dispensed": sum(1 for d in drugs if d["dispensed"]),
            "bills": bl, "billed": round(sum(b["subtotal"] - b["discount"] for b in bl), 2),
            "hmo_share": round(sum(b["hmo_covered"] for b in bl), 2),
            "patient_owes": round(sum(max(b["patient_balance"] - b["paid"], 0) for b in bl), 2),
            "claim_bill_id": bl[0]["id"] if bl else None,
            "claim_status": next((b["claim_status"] for b in bl if b.get("claim_status")), None),
            "approvals_pending": a.get("requested", 0), "approvals": a,
            "signed": bool(e.get("claim_signature_url")) or any(d["signed"] for d in drugs),
            "signed_at": e["claim_signed_at"], "signed_by": e["claim_signed_by"],
            "verified_at": e["hmo_verified_at"], "verified_by": e["hmo_verified_by"], "verify_note": e["hmo_verify_note"],
        })
    return out


@router.get("/hmo-desk")
def hmo_desk(days: int = Query(0, ge=0, le=14, description="also show HMO patients discharged today (0) or N days back"),
             user=Depends(verify_jwt)):
    """The HMO patients on the queue now (and those discharged in the last `days` days)."""
    rows = _hmo_desk_rows(days)
    active = [r for r in rows if r["status"] != "discharged"]
    return {
        "patients": rows,
        "summary": {
            "on_queue": len(active),
            "awaiting_pharmacy": sum(1 for r in active if any(not d["dispensed"] for d in r["drugs"])),
            "to_sign": sum(1 for r in rows if r["bills"] and not r["signed"]),
            "to_verify": sum(1 for r in rows if r["bills"] and not r["verified_at"]),
            "approvals_pending": sum(r["approvals_pending"] for r in rows),
            "hmo_share": round(sum(r["hmo_share"] for r in rows), 2),
        },
    }


@router.get("/hmo-desk/{encounter_id}")
def hmo_desk_patient(encounter_id: int, user=Depends(verify_jwt)):
    rows = _hmo_desk_rows(3650, encounter_id)
    if not rows:
        raise HTTPException(status_code=404, detail="Not an HMO visit")
    return rows[0]


@router.post("/hmo-desk/{encounter_id}/signature")
async def hmo_desk_signature(encounter_id: int, file: UploadFile = File(...), user=Depends(verify_jwt)):
    """The patient signs the visit's claims form on the signature pad; it prints on the form straight away."""
    from src.signatures import ALLOWED_SIGNATURE_TYPES, MAX_SIGNATURE_BYTES, save_signature
    enc = db.table("royan_encounters").select("id").eq("id", encounter_id).limit(1).execute().data
    if not enc:
        raise HTTPException(status_code=404, detail="Visit not found")
    if file.content_type not in ALLOWED_SIGNATURE_TYPES:
        raise HTTPException(status_code=422, detail="The signature must be a PNG, JPEG or WEBP image")
    contents = await file.read()
    if len(contents) > MAX_SIGNATURE_BYTES:
        raise HTTPException(status_code=413, detail="Signature image too large (max 5 MB)")
    url = save_signature(contents, f"claim{encounter_id}", file.filename, file.content_type)
    db.table("royan_encounters").update({
        "claim_signature_url": url, "claim_signed_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "claim_signed_by": user.get("email") or user.get("sub"),
    }).eq("id", encounter_id).execute()
    audit_event("claim_signature_captured", {"encounter_id": encounter_id}, actor_id=user.get("sub"),
                subject_type="encounter", subject_id=str(encounter_id))
    return {"ok": True, "signature_url": url}


class HmoVerifyRequest(BaseModel):
    verified: bool = True
    note: Optional[str] = None


@router.post("/hmo-desk/{encounter_id}/verify")
def hmo_desk_verify(encounter_id: int, payload: HmoVerifyRequest, user=Depends(verify_jwt)):
    """The HMO officer has checked the claims form against the visit (or takes the check back)."""
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    res = db.table("royan_encounters").update({
        "hmo_verified_at": now if payload.verified else None,
        "hmo_verified_by": (user.get("email") or user.get("sub")) if payload.verified else None,
        "hmo_verify_note": payload.note,
    }).eq("id", encounter_id).execute()
    if not res.data:
        raise HTTPException(status_code=404, detail="Visit not found")
    audit_event("claim_verified" if payload.verified else "claim_verification_withdrawn",
                {"encounter_id": encounter_id, "note": payload.note}, actor_id=user.get("sub"),
                subject_type="encounter", subject_id=str(encounter_id))
    return {"ok": True, "verified_at": now if payload.verified else None}
